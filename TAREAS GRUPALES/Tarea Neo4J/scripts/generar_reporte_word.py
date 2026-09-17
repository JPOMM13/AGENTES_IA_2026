from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Reporte_Neo4j_Agente_Cotizador.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_table(table, header_fill="1F4E79"):
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F5F8FC")


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for index, header in enumerate(headers):
        hdr[index].text = header
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = str(value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    style_table(table)
    document.add_paragraph()
    return table


def add_code(document, code):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(code)
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    run.font.size = Pt(9)
    return paragraph


def add_heading(document, text, level=1):
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return heading


def build_document():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.12

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Reporte de Neo4j para agente cotizador de eventos")
    title_run.font.size = Pt(20)
    title_run.font.bold = True

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(
        "Base de conocimiento con instancias y razonamiento simbolico"
    )
    subtitle_run.font.size = Pt(12)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Proyecto: Agente de cotizaciones de eventos | Motor: Neo4j local con Podman")

    add_heading(document, "Resumen ejecutivo", 1)
    document.add_paragraph(
        "Este reporte documenta la implementacion local de una instancia de Neo4j "
        "para modelar una base de conocimiento del agente cotizador de eventos. "
        "La entrega esta alineada con la premisa del ejemplo de clase: no se "
        "presenta solo un diagrama conceptual, sino una red con datos concretos, "
        "relaciones consultables y consultas que explican por que una cotizacion "
        "puede emitirse, quedar pendiente u observarse."
    )

    add_heading(document, "Alcance de la tarea", 1)
    add_table(
        document,
        ["Requisito", "Evidencia en el proyecto"],
        [
            ["Instancia Neo4j local", "Servicio Neo4j levantado con Podman y accesible en localhost:7474."],
            ["Analisis de relaciones", "Relaciones COTIZA con tipos DESCRIBE, REQUIERE, INCUMPLE, CONDICIONA y GENERA_COTIZACION."],
            ["Visualizacion de relaciones", "Consulta MATCH p=()-[:COTIZA]->() RETURN p en Neo4j Browser."],
            ["Reporte documentado", "Este documento describe modelo, consultas, hallazgos y pasos de ejecucion."],
        ],
        widths=[2.1, 4.9],
    )

    add_heading(document, "Modelo de la base de conocimiento", 1)
    document.add_paragraph(
        "Para que la visualizacion sea similar a la del ejemplo del profesor, se uso "
        "un grafo compacto. La relacion principal se llama COTIZA y el significado "
        "del enlace se guarda en la propiedad tipo. Esto permite ver una red limpia "
        "en Neo4j Browser y, al mismo tiempo, conservar informacion simbolica para "
        "el analisis."
    )
    add_table(
        document,
        ["Elemento", "Descripcion", "Ejemplo"],
        [
            ["Cliente", "Solicitante del evento.", "Familia Ramos"],
            ["Evento", "Necesidad a cotizar.", "Boda Ramos"],
            ["Producto", "Servicio o recurso requerido.", "Bar de cocteles"],
            ["Lugar", "Ubicacion del evento.", "Club Chosica"],
            ["Cotizacion", "Resultado comercial generado.", "COT-003 Observada"],
            ["Regla", "Condicion simbolica que afecta la cotizacion.", "Regla: disponibilidad"],
            ["COTIZA", "Relacion principal con propiedad tipo.", "REQUIERE, INCUMPLE, CONDICIONA"],
        ],
        widths=[1.35, 3.4, 2.25],
    )

    document.add_page_break()
    add_heading(document, "Datos cargados y evidencia", 1)
    document.add_paragraph(
        "El archivo recomendado para la visualizacion tipo profesor es "
        "data/cotizador_browser_profesor.cypher. La carga validada genera 22 nodos "
        "y 26 relaciones COTIZA."
    )
    add_table(
        document,
        ["Etiqueta", "Cantidad"],
        [
            ["Producto", "5"],
            ["Cliente", "3"],
            ["Evento", "3"],
            ["Lugar", "3"],
            ["Cotizacion", "3"],
            ["Regla", "3"],
            ["Asesor", "2"],
        ],
        widths=[3.5, 1.2],
    )
    add_table(
        document,
        ["Tipo de relacion COTIZA", "Cantidad"],
        [
            ["REQUIERE", "10"],
            ["DESCRIBE", "3"],
            ["ES_EN_LUGAR", "3"],
            ["GENERA_COTIZACION", "3"],
            ["ASIGNADA_A", "3"],
            ["CONDICIONA", "2"],
            ["INCUMPLE", "1"],
            ["DEBE_CUMPLIR", "1"],
        ],
        widths=[3.5, 1.2],
    )

    add_heading(document, "Consultas de visualizacion", 1)
    document.add_paragraph("Vista general del grafo en Neo4j Browser:")
    add_code(document, "MATCH p=()-[:COTIZA]->()\nRETURN p;")
    document.add_paragraph("Eventos con problemas para cotizar:")
    add_code(
        document,
        'MATCH p=(:Evento)-[:COTIZA]->(:Producto)-[r:COTIZA]->(:Regla)\n'
        'WHERE r.tipo IN ["INCUMPLE", "CONDICIONA"]\n'
        "RETURN p;",
    )
    document.add_paragraph("Explicacion visual de por que Demo Day esta observado:")
    add_code(
        document,
        'MATCH p=(:Cliente {nombre:"Startup Andina"})-[:COTIZA*1..3]->(:Regla)\n'
        "RETURN p;",
    )

    document.add_page_break()
    add_heading(document, "Analisis de relaciones", 1)
    document.add_paragraph(
        "El nodo Evento concentra la decision de cotizacion. Desde un Evento se "
        "puede recorrer hacia los productos requeridos, el lugar, la cotizacion "
        "generada y las reglas que condicionan la viabilidad. Esta estructura "
        "permite explicar la salida del agente con evidencia del grafo."
    )
    add_table(
        document,
        ["Caso", "Ruta simbolica", "Interpretacion"],
        [
            [
                "Boda Ramos",
                "Boda Ramos -> Bar de cocteles -> Regla: disponibilidad",
                "La cotizacion queda pendiente porque un producto solicitado no esta disponible.",
            ],
            [
                "Demo Day",
                "Demo Day -> Club Chosica -> Regla: cobertura",
                "El evento queda observado porque el lugar tiene cobertura condicionada.",
            ],
            [
                "Lanzamiento de producto",
                "ACME Peru -> Lanzamiento -> COT-001 Emitida",
                "El evento cumple condiciones principales y genera cotizacion emitida.",
            ],
        ],
        widths=[1.4, 2.8, 2.8],
    )

    add_heading(document, "Pasos para levantar el proyecto", 1)
    document.add_paragraph("Desde la carpeta del proyecto:")
    add_code(document, "podman compose up -d")
    document.add_paragraph("Cargar la base recomendada para Neo4j Browser:")
    add_code(
        document,
        "podman exec -i neo4j_eventos cypher-shell -u neo4j -p password < "
        "data/cotizador_browser_profesor.cypher",
    )
    document.add_paragraph(
        "Abrir Neo4j Browser en http://localhost:7474/browser/ con usuario neo4j "
        "y clave password. Luego ejecutar la consulta de visualizacion general."
    )

    add_heading(document, "Conclusion", 1)
    document.add_paragraph(
        "La implementacion cumple con la premisa de razonamiento simbolico vista "
        "en clase. La base de conocimiento contiene instancias concretas del "
        "agente cotizador y relaciones que permiten visualizar y explicar la "
        "decision de cotizacion. La relacion COTIZA cumple el rol visual que "
        "YAPEO cumple en el ejemplo del profesor, adaptada al dominio de eventos."
    )

    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
